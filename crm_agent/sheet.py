"""Lecture du CRM « partenariats » (onglets Pipeline, Tâches, Contacts, Notes) et jointure par ID entreprise."""
from __future__ import annotations

import csv
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

TABS = ("Pipeline", "Tâches", "Contacts", "Notes")
DATE_FORMATS = ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y", "%d.%m.%Y")
# Statuts d'une affaire terminée (comparés sans accents ni casse). À adapter à vos statuts.
CLOSED_WORDS = ("signe", "conclu", "refus", "decline", "perdu", "abandon", "termine", "archive", "clos")
DECISION_ROLES = ("directeur", "directrice", "dg", "ceo", "pdg", "president", "responsable", "head", "chef", "fondateur", "drh")


def norm(text: object) -> str:
    text = unicodedata.normalize("NFD", str(text or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", text).strip().lower()


def parse_date(value: str) -> date | None:
    value = (value or "").strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            continue
    return None


@dataclass
class Task:
    id: str
    title: str
    type: str
    priority: str
    due: date | None
    status: str
    owner: str

    @property
    def is_open(self) -> bool:
        return norm(self.status) in ("a faire", "en cours")


@dataclass
class Note:
    date: date | None
    type: str
    text: str
    author: str


@dataclass
class Contact:
    name: str
    role: str
    has_mail: bool
    has_phone: bool

    @property
    def is_decision_maker(self) -> bool:
        return any(w in norm(self.role).split() or w in norm(self.role) for w in DECISION_ROLES)


@dataclass
class Company:
    id: str
    name: str
    domain: str = ""
    sector: str = ""
    added: date | None = None
    status: str = ""
    owner: str = ""
    source: str = ""
    tasks: list[Task] = field(default_factory=list)
    notes: list[Note] = field(default_factory=list)
    contacts: list[Contact] = field(default_factory=list)

    @property
    def is_closed(self) -> bool:
        return any(w in norm(self.status) for w in CLOSED_WORDS)


def _table(rows: list[list[str]], must_have: tuple[str, ...]) -> list[dict[str, str]]:
    """Trouve la ligne d'en-tête (parmi les 6 premières) et renvoie des dicts {en-tête normalisé: valeur}.
    En cas d'en-têtes dupliqués (vue de recherche à droite du Pipeline), la 1re occurrence gagne."""
    for h_idx, row in enumerate(rows[:6]):
        headers = [norm(c) for c in row]
        if all(m in headers for m in must_have):
            break
    else:
        raise ValueError(f"En-tête introuvable (colonnes attendues : {must_have})")
    out = []
    for raw in rows[h_idx + 1:]:
        rec: dict[str, str] = {}
        for h, cell in zip(headers, raw):
            if h and h not in rec:
                rec[h] = str(cell).strip()
        if rec.get(must_have[0]):
            out.append(rec)
    return out


def build_companies(tabs: dict[str, list[list[str]]]) -> list[Company]:
    companies = {}
    for r in _table(tabs["Pipeline"], ("id", "nom de l'entreprise", "statut")):
        if not r["id"].upper().startswith("ENT"):
            continue
        companies[r["id"]] = Company(
            id=r["id"], name=r.get("nom de l'entreprise", ""), domain=r.get("nom de domaine", ""),
            sector=r.get("secteur - domaine", ""), added=parse_date(r.get("date d'ajout", "")),
            status=r.get("statut", ""), owner=r.get("personne en charge", ""), source=r.get("source", ""),
        )

    # Les en-têtes de l'onglet Tâches contiennent parfois une coquille sur la colonne ID : on s'appuie sur « id entreprise ».
    for r in _table(tabs["Tâches"], ("id entreprise", "statut")):
        c = companies.get(r["id entreprise"])
        if c:
            first_col = next(iter(r.values()))
            c.tasks.append(Task(first_col, r.get("tache", ""), r.get("type", ""), r.get("priorite", ""),
                                parse_date(r.get("echeance", "")), r.get("statut", ""), r.get("personne en charge", "")))
    for r in _table(tabs["Contacts"], ("id entreprise", "nom du contact")):
        c = companies.get(r["id entreprise"])
        if c:
            c.contacts.append(Contact(r["nom du contact"], r.get("role", ""), bool(r.get("mail")), bool(r.get("telephone"))))
    for r in _table(tabs["Notes"], ("id entreprise", "note")):
        c = companies.get(r["id entreprise"])
        if c:
            c.notes.append(Note(parse_date(r.get("date", "")), r.get("type", ""), r.get("note", ""), r.get("auteur", "")))
    for c in companies.values():
        c.notes.sort(key=lambda n: n.date or date.min)
    return list(companies.values())


# --- Sources -----------------------------------------------------------------------------------

def _cell(v: object) -> str:
    if isinstance(v, datetime):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, date):
        return v.strftime("%d/%m/%Y")
    return "" if v is None else str(v)


def load_xlsx(path: str) -> list[Company]:
    """Export du Sheet : Fichier > Télécharger > Microsoft Excel (.xlsx)."""
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    tabs = {name: [[_cell(c) for c in row] for row in wb[name].iter_rows(values_only=True)] for name in TABS}
    return build_companies(tabs)


def load_csv_dir(directory: str) -> list[Company]:
    """Dossier contenant Pipeline.csv, Tâches.csv, Contacts.csv, Notes.csv."""
    tabs = {}
    for name in TABS:
        with open(Path(directory) / f"{name}.csv", newline="", encoding="utf-8-sig") as f:
            tabs[name] = [list(r) for r in csv.reader(f)]
    return build_companies(tabs)


def load_google_sheet(spreadsheet: str, credentials_path: str | None) -> list[Company]:
    """`spreadsheet` = URL ou ID. Auth par compte de service (la feuille doit lui être partagée)."""
    import gspread

    gc = gspread.service_account(filename=credentials_path) if credentials_path else gspread.service_account()
    book = gc.open_by_url(spreadsheet) if spreadsheet.startswith("http") else gc.open_by_key(spreadsheet)
    return build_companies({name: book.worksheet(name).get_all_values() for name in TABS})
